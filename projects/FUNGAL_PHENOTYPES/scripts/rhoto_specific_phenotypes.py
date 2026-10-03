#!/usr/bin/env python3
"""Find condition-specific fitness defects in Rhodotorula toruloides RB-TDNAseq.

Reads the gene fitness matrix of Kim et al. 2021 (PMID:33585414; Frontiers
supplement Table 2, sheet "RB-TDNA Seq", which also carries the Coradetti et
al. 2018 conditions) and the RTO4 -> UniProt mapping of Coradetti et al. 2023
(PMID:37537586; MOESM1, sheet "ProteinInfo"), then calls a gene/condition
pair a *specific defect* when

* fitness in the condition is <= --strong (default -2), and
* the gene is not generally sick: fitness in every glucose control is
  > --sick (default -1).

The matrix has fitness scores only (no t statistics), so the thresholds are
deliberately stricter than the |fit| > 1 used with t-values in the Fitness
Browser. Each defective gene is annotated with its current UniProtKB status
(many IFO0880 accessions have since been deleted and survive only in UniParc),
protein name and annotation score.

Used by ``projects/FUNGAL_PHENOTYPES/RHOTO.md``. Nothing is hardcoded:

    uv run python projects/FUNGAL_PHENOTYPES/scripts/rhoto_specific_phenotypes.py

Downloads are cached in ``--cache`` (default ``tmp/fungal_phenotypes``).
"""

from __future__ import annotations

import argparse
import collections
import csv
import io
import json
import sys
import time
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

import openpyxl

EPMC_SUPP = "https://www.ebi.ac.uk/europepmc/webservices/rest/{}/supplementaryFiles"
KIM_PMC, KIM_FILE = "PMC7873862", "Table_2.XLSX"
CORADETTI_PMC, CORADETTI_FILE = "PMC10398944", "12934_2023_2148_MOESM1_ESM.xlsx"
UNIPROT_SEARCH = "https://rest.uniprot.org/uniprotkb/search"

GLUCOSE_CONTROLS = ["M9_Glucose", "YNB Glucose", "YNB_CSM_KPO4 Glucose"]
# Conditions that are not a change of carbon source: glucose media with
# supplements, and nitrogen starvation (lipid mobilisation).
NON_CARBON = GLUCOSE_CONTROLS + [
    "YNB Glucose plus Arginine",
    "YNB Glucose plus Methionine",
    "YNB Glucose plus Dropout Complete",
    "Fitness During Lipid Mobilization",
]


def supplement(pmc: str, name: str, cache: Path) -> Path:
    dest = cache / f"{pmc}_{name}"
    if not dest.exists():
        cache.mkdir(parents=True, exist_ok=True)
        print(f"downloading {pmc} supplementary files", file=sys.stderr)
        with urllib.request.urlopen(EPMC_SUPP.format(pmc), timeout=600) as r:
            blob = r.read()
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            dest.write_bytes(z.read(name))
    return dest


def read_fitness(path: Path) -> tuple[list[str], dict[int, dict[str, float]]]:
    ws = openpyxl.load_workbook(path, read_only=True)["RB-TDNA Seq"]
    rows = ws.iter_rows(values_only=True)
    header = [h for h in next(rows) if h]
    conds = header[1:]
    fit = {}
    for row in rows:
        vals = dict(zip(conds, row[1 : len(header)]))
        if any(v is not None for v in vals.values()):
            fit[int(row[0])] = {c: v for c, v in vals.items() if v is not None}
    return conds, fit


def read_mapping(path: Path) -> dict[int, dict[str, str]]:
    ws = openpyxl.load_workbook(path, read_only=True)["ProteinInfo"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    out = {}
    for row in rows:
        rec = dict(zip(header, row))
        out[int(rec["Protein"])] = {
            "uniprot": rec["Uniprot"] if rec["Uniprot"] not in (None, "#N/A") else "",
            "annotation": rec["Combined Annotations"] or "",
            "sc_orthologs": rec["Sc288c Orthologs"] or "",
            "human_orthologs": rec["Human Orthologs"] or "",
        }
    return out


def uniprot_status(accs: list[str], cache: Path) -> dict[str, dict[str, str]]:
    """Current UniProtKB name / annotation score, or 'deleted'."""
    path = cache / "rhoto_uniprot_status.json"
    known = json.loads(path.read_text()) if path.exists() else {}
    todo = [a for a in accs if a and a not in known]
    for i in range(0, len(todo), 100):
        batch = todo[i : i + 100]
        params = {
            "query": " OR ".join(f"accession:{a}" for a in batch),
            "fields": "accession,protein_name,annotation_score,reviewed",
            "format": "tsv",
            "size": 500,
        }
        url = f"{UNIPROT_SEARCH}?{urllib.parse.urlencode(params)}"
        with urllib.request.urlopen(url, timeout=120) as r:
            lines = r.read().decode().splitlines()[1:]
        for line in lines:
            f = line.split("\t") + [""] * 4
            known[f[0]] = {
                "status": "deleted" if f[1] == "deleted" else f[3] or "active",
                "protein_name": "" if f[1] == "deleted" else f[1],
                "annotation_score": f[2],
            }
        for a in batch:
            known.setdefault(a, {"status": "not_found", "protein_name": "", "annotation_score": ""})
        time.sleep(0.2)
    path.write_text(json.dumps(known, indent=0))
    return known


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--cache", type=Path, default=Path("tmp/fungal_phenotypes"))
    ap.add_argument("--strong", type=float, default=-2.0)
    ap.add_argument("--sick", type=float, default=-1.0)
    ap.add_argument(
        "--outdir", type=Path, default=Path("projects/FUNGAL_PHENOTYPES/data")
    )
    args = ap.parse_args()

    conds, fit = read_fitness(supplement(KIM_PMC, KIM_FILE, args.cache))
    mapping = read_mapping(supplement(CORADETTI_PMC, CORADETTI_FILE, args.cache))
    carbon = [c for c in conds if c not in NON_CARBON]

    hits = collections.defaultdict(dict)  # gene -> {condition: fitness}
    n_sick = 0
    for gene, f in fit.items():
        controls = [f[c] for c in GLUCOSE_CONTROLS if c in f]
        if not controls or min(controls) <= args.sick:
            n_sick += bool(controls)
            continue
        for c in carbon:
            if f.get(c, 0) <= args.strong:
                hits[gene][c] = f[c]

    accs = sorted({mapping.get(g, {}).get("uniprot", "") for g in hits} - {""})
    status = uniprot_status(accs, args.cache)

    args.outdir.mkdir(parents=True, exist_ok=True)
    out = args.outdir / "rhoto_specific_defects.tsv"
    with out.open("w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t", lineterminator="\n")
        w.writerow(
            [
                "rto4_id",
                "n_carbon_conditions_defective",
                "defective_conditions",
                "min_glucose_control_fitness",
                "uniprot_2023",
                "uniprot_status",
                "uniprot_protein_name",
                "uniprot_annotation_score",
                "kog_annotation",
                "sc_orthologs",
                "human_orthologs",
            ]
        )
        for gene in sorted(hits, key=lambda g: (len(hits[g]), min(hits[g].values()))):
            m = mapping.get(gene, {})
            s = status.get(m.get("uniprot", ""), {})
            f = fit[gene]
            w.writerow(
                [
                    f"RTO4_{gene}",
                    len(hits[gene]),
                    "; ".join(
                        f"{c}={v:.1f}" for c, v in sorted(hits[gene].items(), key=lambda x: x[1])
                    ),
                    f"{min(f[c] for c in GLUCOSE_CONTROLS if c in f):.2f}",
                    m.get("uniprot", ""),
                    s.get("status", ""),
                    s.get("protein_name", ""),
                    s.get("annotation_score", ""),
                    m.get("annotation", ""),
                    m.get("sc_orthologs", ""),
                    m.get("human_orthologs", ""),
                ]
            )

    per_cond = collections.Counter(c for h in hits.values() for c in h)
    st = collections.Counter(status.get(mapping[g]["uniprot"], {}).get("status") for g in hits)
    print(
        f"{len(fit)} genes with fitness; {n_sick} excluded as sick on glucose; "
        f"{len(hits)} genes with >=1 specific carbon-source defect",
        file=sys.stderr,
    )
    print(f"UniProtKB status of their 2023 accessions: {dict(st)}", file=sys.stderr)
    for c in carbon:
        print(f"  {c}\t{per_cond[c]}", file=sys.stderr)
    print(f"wrote {out}", file=sys.stderr)


if __name__ == "__main__":
    main()
