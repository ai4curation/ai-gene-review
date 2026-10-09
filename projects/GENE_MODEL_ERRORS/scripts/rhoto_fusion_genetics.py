#!/usr/bin/env python3
"""Test R. toruloides fusion candidates with transposon fitness data.

For each target called ``FUSION`` in ``data/rhoto_fusion_screen.tsv``
(``tile_gene_model.py`` against the IFO0880 proteome), maps each tiling
component to its IFO0880 gene (RTO4 id, via the Coradetti et al. 2023
ProteinInfo sheet) and its RB-TDNAseq fitness profile (Kim et al. 2021), and
records two independent lines of evidence:

* ``adjacent`` -- the components are consecutive RTO4 ids, i.e. neighbouring
  genes in the IFO0880 annotation;
* ``genetically_separated`` -- in some condition one component's mutants have
  a strong defect (<= --strong) while another component's mutants grow
  normally (> --normal). Mutating one half of a single protein would be
  expected to affect the whole protein, so independent phenotypes show the
  halves are separate genes.

    uv run python projects/GENE_MODEL_ERRORS/scripts/rhoto_fusion_genetics.py

Reads the supplements cached by
``projects/FUNGAL_PHENOTYPES/scripts/rhoto_specific_phenotypes.py``.
"""

from __future__ import annotations

import argparse
import collections
import csv
import sys
from pathlib import Path

import openpyxl

DATA = Path("projects/GENE_MODEL_ERRORS/data")
CACHE = Path("tmp/fungal_phenotypes")
KIM = CACHE / "PMC7873862_Table_2.XLSX"
CORADETTI = CACHE / "PMC10398944_12934_2023_2148_MOESM1_ESM.xlsx"


def rto4_by_accession() -> dict[str, str]:
    ws = openpyxl.load_workbook(CORADETTI, read_only=True)["ProteinInfo"]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    return {r[2]: r[1] for r in rows if r[2]}


def fitness() -> dict[str, dict[str, float]]:
    ws = openpyxl.load_workbook(KIM, read_only=True)["RB-TDNA Seq"]
    rows = ws.iter_rows(values_only=True)
    header = [h for h in next(rows) if h]
    out = {}
    for r in rows:
        vals = {c: v for c, v in zip(header[1:], r[1 : len(header)]) if v is not None}
        if vals:
            out[f"RTO4_{r[0]}"] = vals
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--strong", type=float, default=-2.5)
    ap.add_argument("--normal", type=float, default=-1.0)
    args = ap.parse_args()

    if not (KIM.exists() and CORADETTI.exists()):
        sys.exit("run projects/FUNGAL_PHENOTYPES/scripts/rhoto_specific_phenotypes.py first")
    rto4 = rto4_by_accession()
    fit = fitness()

    targets = collections.defaultdict(list)
    for r in csv.DictReader((DATA / "rhoto_fusion_screen.tsv").open(), delimiter="\t"):
        if r["call"] == "FUSION":
            targets[r["target"]].append(r)

    out = DATA / "rhoto_fusion_genetics.tsv"
    n_sep = n_adj = 0
    with out.open("w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t", lineterminator="\n")
        w.writerow(
            ["target", "target_length", "components", "adjacent",
             "genetically_separated", "separating_condition", "evidence"]
        )
        for target, segs in targets.items():
            comps = []
            for s in segs:
                accs = [a.split(".")[0] for a in s["comparison_accessions"].split("; ")]
                gene = next((rto4[a] for a in accs if a in rto4), "")
                comps.append((s, gene))
            genes = [g for _, g in comps if g]
            nums = sorted(int(g.split("_")[1]) for g in genes)
            adjacent = len(nums) == len(comps) and nums == list(range(nums[0], nums[0] + len(nums)))
            sep, why = False, ""
            for _, ga in comps:
                for _, gb in comps:
                    if not ga or not gb or ga == gb or ga not in fit or gb not in fit:
                        continue
                    for cond, va in fit[ga].items():
                        vb = fit[gb].get(cond)
                        if va <= args.strong and vb is not None and vb > args.normal:
                            if not sep or va < float(why.split("=")[1].split(",")[0]):
                                sep, why = True, f"{cond}: {ga}={va:.1f}, {gb}={vb:.1f}"
            n_sep += sep
            n_adj += adjacent
            evidence = ["TILING"] + (["ADJACENT_GENES"] if adjacent else []) + (
                ["GENETIC_SEPARATION"] if sep else []
            )
            w.writerow([
                target, segs[0]["target_length"],
                "; ".join(f"{s['target_from']}-{s['target_to']}={g or '?'}" for s, g in comps),
                adjacent, sep, why, ",".join(evidence),
            ])
    print(
        f"{len(targets)} fusions: {n_adj} with adjacent components, "
        f"{n_sep} genetically separated; wrote {out}",
        file=sys.stderr,
    )


if __name__ == "__main__":
    main()
