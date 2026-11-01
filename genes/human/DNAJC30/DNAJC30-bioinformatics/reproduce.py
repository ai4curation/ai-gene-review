# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1,<4"]
# ///
"""Reproduce the factual DNAJC30 Table S1 extract from an unedited public workbook.

Use --workbook for saved-only operation, or --download for the publisher URL.
The expected source hash is in source-manifest.json. No measurement is hardcoded.
"""
import argparse
import hashlib
import io
import json
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

HERE = Path(__file__).resolve().parent


def extract(data, source):
    digest = hashlib.sha256(data).hexdigest()
    if digest != source["sha256"] or len(data) != source["bytes"]:
        raise ValueError("Source bytes differ from the recorded original workbook; do not silently substitute a new version")
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    # Coordinate selection follows Table S1's layout, not an inferred biological result.
    specs = [
        ("(A) All protein groups", "M", "C", "Gene names", "Protein IDs",
         "A B C D E F G H M Q AF AG AH AI AJ AK AL AM AN AO AP BN BO BP BQ BR BS BT BU BV BW BX BY BZ CA CB CC CD CE CF CG CH CI CJ CK CL".split()),
        ("(B) MitoCoP (1,134 genes)", "C", "B", "Gene name", "Simplified protein IDs", "A B C D".split()),
    ]
    records = []
    for name, gene_col, accession_col, gene_header, accession_header, columns in specs:
        sheet = workbook[name]
        rows = sheet.iter_rows(values_only=True)
        next(rows)  # group-label row
        headers = next(rows)
        gi, ai = column_index_from_string(gene_col)-1, column_index_from_string(accession_col)-1
        if headers[gi] != gene_header or headers[ai] != accession_header:
            raise ValueError(f"Unexpected identity columns in {name}")
        matches = []
        for excel_row, row in enumerate(rows, 3):
            genes = str(row[gi] or "").split(";")
            accessions = str(row[ai] or "").split(";")
            if source["target"]["gene"] in genes and source["target"]["accession"] in accessions:
                fields = []
                for column in columns:
                    i = column_index_from_string(column)-1
                    fields.append({"cell": f"{column}{excel_row}", "header": headers[i], "value": row[i]})
                matches.append({"sheet": name, "row_one_based": excel_row, "fields": fields})
        if len(matches) != 1:
            raise ValueError(f"Expected one exact target row in {name}; found {len(matches)}")
        records.extend(matches)
    workbook.close()
    return {"reference_id": source["reference_id"], "table": source["table"], "source_url": source["url"],
            "source_sha256": digest, "target": source["target"], "scope": source["scope"], "records": records}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--workbook", type=Path)
    inputs.add_argument("--download", action="store_true")
    outputs = parser.add_mutually_exclusive_group(required=True)
    outputs.add_argument("--output", type=Path)
    outputs.add_argument("--check", type=Path)
    args = parser.parse_args()
    source = json.loads((HERE / "source-manifest.json").read_text())
    if args.workbook:
        data = args.workbook.read_bytes()
    else:
        with urlopen(Request(source["url"], headers={"User-Agent": "DNAJC30-source-reproduction/1.0"}), timeout=90) as response:
            data = response.read()
    result = extract(data, source)
    if args.check:
        if result != json.loads(args.check.read_text()):
            raise ValueError("Regenerated measurements differ from the supplied extract")
        print(f"PASS: {len(result['records'])} target rows; {sum(len(r['fields']) for r in result['records'])} original workbook cells")
    else:
        args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False)+"\n")
        print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
