"""Reproduce EFNB1 evidence from hash-pinned saved data; no network access."""
from pathlib import Path
import base64, csv, hashlib, json, re, zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from bs4 import BeautifulSoup

C = Path(__file__).resolve().parent
O = C
S = C
D = C / 'projections'
R = None
PROJECTIONS = json.loads((C / 'source-projections.json').read_text())
def J(path):
    return PROJECTIONS[path.name] if path.parent == D else json.loads(path.read_text())
H = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
json_value = lambda value: json.loads(json.dumps(value))
def require(value, message):
    if not value:
        raise ValueError(message)

MITAB = {0:'id_a',1:'id_b',6:'method',7:'author',8:'publication',9:'taxon_a',10:'taxon_b',11:'interaction_type',13:'interaction_id',14:'confidence',15:'expansion',16:'biological_role_a',17:'biological_role_b',18:'experimental_role_a',19:'experimental_role_b',27:'annotations',28:'host',32:'checksum_a',33:'checksum_b',35:'negative',36:'feature_a',37:'feature_b',40:'identification_a',41:'identification_b'}

def pairs():
    rows = list(csv.reader((R/'intact-P98172.tsv').open(), delimiter='\t'))
    selected = []
    for number, cells in enumerate(rows, 1):
        require(len(cells) >= 42, f'Short MITAB line {number}')
        pmids = cells[8].split('|')
        partner_ids = {cells[0], cells[1]}
        take = any('pubmed:'+p in pmids and partner_ids == {'uniprotkb:P98172','uniprotkb:'+partner} for p,partner in (('9883737','Q62083'),('22279592','P04626'),('32296183','Q86W74-2')))
        take |= ('pubmed:33961781' in pmids and partner_ids in ({'uniprotkb:P98172','uniprotkb:P29323'},{'uniprotkb:P98172','uniprotkb:P04626'}))
        if take and 'uniprotkb:P98172' in partner_ids:
            selected.append({'physical_line':number, **{name:cells[i] for i,name in MITAB.items()}})
    checks = []
    for record in selected:
        for side in ('a','b'):
            accession = record['id_'+side].split(':',1)[1]
            base = accession.split('-')[0]
            path = R/f'UniProt-{base}.json'
            official = J(path)
            sequence = official['sequence']['value']
            if '-' in accession:
                products = [x for c in official.get('comments',[]) if c['commentType']=='ALTERNATIVE PRODUCTS' for x in c.get('isoforms',[]) if accession in x['isoformIds']]
                require(len(products)==1 and products[0]['isoformSequenceStatus']=='Displayed', f'Unresolved isoform {accession}')
                fasta = ''.join(l for l in (R/f'{accession}.fasta').read_text().splitlines() if not l.startswith('>'))
                require(fasta == sequence, 'Displayed isoform FASTA mismatch')
            taxon = str(official['organism']['taxonId'])
            calculated = 'rogid:'+base64.b64encode(hashlib.sha1(sequence.encode()).digest()).decode().rstrip('=')+taxon
            observed = record['checksum_'+side]
            require(calculated == observed, f'Reference checksum mismatch {record["physical_line"]} {side}')
            checks.append({'line':record['physical_line'],'side':side,'accession':accession,'official_path':path.name,'length':len(sequence),'observed':observed,'calculated':calculated,'match':True,'meaning':'Curated reference sequence checksum; assayed fragment/tag is a separate feature field.'})
    expected = J(D/'selected-pair-reference-proof-v1.json')
    require(selected == expected['records'], 'Selected source records differ')
    # The original proof used repository-relative paths, preserved without changing bytes.
    for actual, saved in zip(checks, expected['sequence_checks']):
        require({k:v for k,v in actual.items() if k!='official_path'} == {k:v for k,v in saved.items() if k!='official_path'}, 'Sequence proof changed')
        require(Path(saved['official_path']).name == Path(actual['official_path']).name, 'Official reference identity')
    require(len(checks)==len(expected['sequence_checks']), 'Sequence-check count')
    return {'all_records':len(rows),'selected':len(selected),'reference_checks':len(checks)}

def reactome():
    expected = J(D/'reactome-target-proof-v2.json')
    nodes = {}
    for p in sorted(R.glob('Reactome*.json')):
        if '.receipt.' in p.name:
            continue
        try:
            obj = J(p)
        except (ValueError,UnicodeDecodeError):
            continue
        if isinstance(obj,dict) and 'dbId' in obj:
            nodes[obj['dbId']] = obj
    for key, filename in expected['files'].items():
        nodes[int(key)] = J(R/filename)
    relations = ['input','output','hasComponent','hasMember','hasCandidate','physicalEntity','catalystActivity','repeatedUnit']
    def walk(i,path):
        if i in path:
            return []
        require(i in nodes, f'Missing saved Reactome node {i}')
        node = nodes[i]; path = path+[i]; ref = node.get('referenceEntity',{})
        if isinstance(ref,dict) and ref.get('identifier')=='P98172':
            return [{'path':path,'referenceEntity':ref,'compartments':node.get('compartment',[]),'coordinates':{k:node.get(k) for k in ('startCoordinate','endCoordinate')}}]
        hits = []
        for relation in relations:
            children = node.get(relation,[])
            children = children if isinstance(children,list) else [children]
            for child in children:
                ci = child.get('dbId') if isinstance(child,dict) else child
                if not ci:
                    continue
                name = child.get('displayName','') if isinstance(child,dict) else nodes.get(ci,{}).get('displayName','')
                if name and not any(t in name.lower() for t in ['efn','ephrin']) and relation!='catalystActivity':
                    continue
                for hit in walk(ci,path):
                    hit.setdefault('relations',[]).insert(0,relation)
                    hits.append(hit)
        return hits
    rows = []
    for i, source in enumerate(J(C/'source-assertions.json')):
        ref = source['original_reference_id']
        if ref.startswith('Reactome:'):
            hits = walk(int(ref.rsplit('-',1)[1]),[])
            require(hits, f'Target absent in saved trace {ref}')
            require(all(h['referenceEntity']['identifier']=='P98172' and any(c.get('accession')=='0005886' for c in h['compartments']) for h in hits), 'Target location mismatch')
            rows.append({'source_index':i,'reference':ref,'hits':hits})
    require(rows == expected['source_rows'], 'Reactome graph replay changed')
    return {'source_rows':len(rows),'positive_paths':sum(len(x['hits']) for x in rows),'saved_used_nodes':len(expected['files'])}

def paint():
    expected = J(D/'paint-proof-v1.json')
    root = J(R/'PTHR11304-tree.json')['search']['tree_topology']['annotation_node']
    nodes,paths = {},{}
    def walk(node,path):
        key=node['persistent_id'];nodes[key]=node;paths[key]=path+[key]
        children=node.get('children',{}).get('annotation_node',[])
        if isinstance(children,dict):children=[children]
        for child in children:walk(child,path+[key])
    walk(root,[])
    target=[n for n in nodes.values() if n.get('node_name','').endswith('UniProtKB=P98172')]
    require(len(target)==1,'Ambiguous PAINT target')
    key=target[0]['persistent_id']
    align=[x for x in J(R/'PTHR11304-msa.json')['search']['MSA_list']['sequence_info'] if x['persistent_id']==key]
    require(len(align)==1,'Ambiguous target MSA')
    sequence=re.sub('[^A-Za-z]','',align[0]['sequence']).upper()
    require(sequence==J(R/'UniProt-P98172.json')['sequence']['value'],'Target alignment mismatch')
    require(len(nodes)==expected['tree_nodes'] and key==expected['target_node'] and paths[key]==expected['path'],'Tree placement mismatch')
    ibd=R/expected['IBD_path'];require(H(ibd)==expected['IBD_sha256'],'Changed IBD source')
    selected=[]
    for line,text in enumerate(ibd.open(),1):
        cells=text.rstrip('\n').split('\t')
        if len(cells)>6 and cells[1] in ('PTN000885159','PTN002601663'):
            selected.append({'physical_line':line,'fields':cells,'node_in_target_path':cells[1] in paths[key]})
    require(selected==expected['IBD_rows'] and all(x['node_in_target_path'] for x in selected),'Literal IBD placement mismatch')
    for accession,saved in expected['current_donor_experimental_annotations'].items():
        source=J(R/f'QuickGO-{accession}-page1.json')
        require(len(source['results'])==source['numberOfHits'],'Incomplete donor query')
        require([x for x in source['results'] if x['goEvidence'] in ('IDA','IMP','IPI','IGI','IEP')]==saved,'Donor projection mismatch')
    return {'nodes':len(nodes),'target_length':len(sequence),'literal_IBDs':len(selected)}

def measurements():
    expected=J(D/'measurement-proof-v1.json')
    for p in expected['input_pins']:
        require(H(R/p['path'])==p['sha256'],'Changed measurement input '+p['path'])
    # Extracted workbook bytes are individually authenticated; the parent archive is not redistributed.
    with (R/'35922511-41586_2022_5028_MOESM4_ESM.xlsx').open('rb') as stream:
        workbook=load_workbook(stream,read_only=True,data_only=True)
        rows=list(workbook['Table 1'].values); hits=[(i+1,r) for i,r in enumerate(rows) if r[0]=='P98172'];require(len(hits)==1,'Construct ambiguity')
        number,row=hits[0]; construct={str(k):v for k,v in zip(rows[5],row) if k};construct['excel_row']=number
        sequence=J(R/'UniProt-P98172.json')['sequence']['value']
        construct.update(sequence_length=len(construct['AA_seq']),matches_official_prefix=sequence.startswith(construct['AA_seq']))
        require(construct==expected['35922511']['library_construct'],'Construct projection mismatch')
        rows=list(workbook['Table 2'].values); headers=rows[5];matrix=[]
        for partner in ('EPHA2','EPHB2','EPHA4'):
            for rname,cname in (('EFNB1',partner),(partner,'EFNB1')):
                ix=[i for i,r in enumerate(rows) if r[0]==rname];jx=[j for j,v in enumerate(headers) if v==cname];require(len(ix)==len(jx)==1,'Matrix target ambiguity')
                i,j=ix[0],jx[0];matrix.append({'row_label':rname,'column_label':cname,'cell':f'{get_column_letter(j+1)}{i+1}','raw_absorbance':rows[i][j]})
        require(matrix==expected['35922511']['primary_raw_matrix'],'Matrix values differ')
        rows=list(workbook['Table 3'].values);literature=[{'excel_row':i+1,'fields':dict(zip(rows[6],r))}for i,r in enumerate(rows) if 'P98172' in r]
        require(json_value(literature)==expected['35922511']['curated_literature_not_new_assay'],'Literature projection mismatch')
        workbook.close()
    workbook=load_workbook(R/'35922511-41586_2022_5028_MOESM6_ESM.xlsx',read_only=True,data_only=True)
    for sheet,saved in expected['35922511']['source_data'].items():
        rows=list(workbook[sheet].values);incident=[(i+1,r)for i,r in enumerate(rows[1:],1)if 'EFNB1' in r or 'P98172' in r]
        selected=[{'excel_row':i,'fields':dict(zip(rows[0],r))}for i,r in incident if any(p in r for p in ('EPHA2','EPHB2','EPHA4'))]
        require(json_value({'all_data_rows':len(rows)-1,'incident_EFNB1_count':len(incident),'selected_three_partner_rows':selected})==saved,'Source data mismatch '+sheet)
    workbook.close()
    graph_path=next(R/x['path']for x in expected['input_pins']if Path(x['path']).name=='40205054-apms-network.json')
    graph=J(graph_path);cx={k:[v for part in graph for v in part.get(k,[])]for k in ('nodes','edges','nodeAttributes','edgeAttributes','networkAttributes')};nodes={x['@id']:x for x in cx['nodes']}
    target=[k for k,v in nodes.items()if v.get('n')=='EFNB1'];require(len(target)==1,'Graph target ambiguity');tid=target[0]
    incident=[e for e in cx['edges']if tid in (e['s'],e['t'])];selected=[]
    for edge in incident:
        other=nodes[edge['t']if edge['s']==tid else edge['s']]
        if other['n']=='EPHB2':selected.append({'edge':edge,'source_node':nodes[edge['s']],'target_node':nodes[edge['t']],'edge_attributes':[a for a in cx['edgeAttributes']if a['po']==edge['@id']],'node_attributes':[a for a in cx['nodeAttributes']if a['po']in(edge['s'],edge['t'])]})
    require({'nodes':len(nodes),'edges':len(cx['edges']),'incident_EFNB1_edges':len(incident),'selected':selected,'network_attributes':cx['networkAttributes']}==expected['40205054'],'Graph projection mismatch')
    html=BeautifulSoup((R/'19056867-database.html').read_bytes(),'html.parser');out=[];tables=html.find_all('table')
    for ti,table in enumerate(tables):
        rows=table.find_all('tr')
        for i,row in enumerate(rows):
            cells=[c.get_text(' ',strip=True)for c in row.find_all(['th','td'])]
            if 'EFNB1' in cells:out.append({'html_table_index':ti,'table_row_zero_based':i,'cells':cells,'preceding_header_rows':[[c.get_text(' ',strip=True)for c in r.find_all(['th','td'])]for r in rows[:3]],'links':[a.get('href')for a in row.find_all('a')]})
    require(out==expected['19056867'],'Exosome table projection mismatch')
    return {'human_construct_residues':construct['sequence_length'],'raw_matrix_orientations':len(matrix),'APMS_incident_edges':len(incident),'APMS_selected_edges':len(selected),'exosome_rows':len(out)}

def run():
    return {'pairs':pairs(),'reactome':reactome(),'paint':paint(),'measurements':measurements()}

def main():
    import argparse, tempfile
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inputs',type=Path,default=C/'inputs.zip',help='Exact saved input ZIP; no download is attempted')
    parser.add_argument('--output',type=Path,help='Optional JSON summary output')
    args=parser.parse_args()
    manifest=json.loads((C/'source-manifest.json').read_text())
    require(H(args.inputs)==manifest['archive']['sha256'],'Input archive hash mismatch')
    with tempfile.TemporaryDirectory(prefix='efnb1-replay-') as temporary:
        global R
        R=Path(temporary)
        with zipfile.ZipFile(args.inputs) as archive:
            require(set(archive.namelist())=={x['path']for x in manifest['inputs']},'Input member set differs')
            for item in manifest['inputs']:
                name=item['path']
                require(Path(name).name==name,'Non-flat archive path')
                data=archive.read(name)
                require(len(data)==item['bytes'] and hashlib.sha256(data).hexdigest()==item['sha256'],'Input hash mismatch '+name)
                (R/name).write_bytes(data)
        summary=run()
        hpa=BeautifulSoup((R/'HPA-P98172-correct-subcellular.html').read_bytes(),'html.parser').get_text(' ',strip=True)
        require('Plasma membrane (supported)'in hpa and 'ENSG00000090776'in(R/'HPA-P98172-correct.xml').read_text(),'HPA target or reliability changed')
        result={'status':'PASS','input_files':len(manifest['inputs']),'summary':summary,'scope':'Reference-sequence identities and raw recorded measurements are reproduced. They do not establish native interfaces, unique peptide identity, or unmeasured activities.'}
        text=json.dumps(result,indent=2)+'\n'
        if args.output:
            args.output.write_text(text)
        else:
            print(text,end='')
if __name__=='__main__':
    main()
