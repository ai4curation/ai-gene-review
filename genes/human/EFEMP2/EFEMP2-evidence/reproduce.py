#!/usr/bin/env python3
"""Reconstruct selected database records from exact, hash-checked source files.

This extracts evidence; it does not assign annotation actions. A reference-sequence
checksum is not authentication of an experimental insert or fragment.
"""
import argparse
import base64
import csv
import hashlib
import json
from pathlib import Path
import urllib.request


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def load_inputs(manifest, mapping, download):
    paths = {}
    for key, spec in manifest['inputs'].items():
        if key in mapping:
            p = Path(mapping[key])
        elif download:
            p = download / spec['filename']
            if not p.exists():
                download.mkdir(parents=True, exist_ok=True)
                with urllib.request.urlopen(spec['url'], timeout=120) as source:
                    data = source.read()
                if hashlib.sha256(data).hexdigest() != spec['sha256']:
                    raise ValueError('Downloaded version differs from recorded source: ' + key)
                p.write_bytes(data)
        else:
            raise ValueError('No saved input for ' + key + '; use --download-cache if appropriate')
        if digest(p) != spec['sha256']:
            raise ValueError('Source hash mismatch: ' + key)
        paths[key] = p
    return paths


def official_sequences(paths):
    records = {k[8:]: json.loads(p.read_text()) for k, p in paths.items() if k.startswith('uniprot:')}
    seqs = {k: v['sequence']['value'] for k, v in records.items()}
    variants = []
    for key, path in paths.items():
        if not key.startswith('fasta:'):
            continue
        acc = key[6:]
        parent = records[acc.split('-')[0]]
        choices = [i for c in parent.get('comments', []) if c['commentType'] == 'ALTERNATIVE PRODUCTS'
                   for i in c['isoforms'] if acc in i['isoformIds']]
        if len(choices) != 1:
            raise ValueError('Nonunique official product: ' + acc)
        iso = choices[0]
        seq = parent['sequence']['value']
        edits = [f for f in parent['features'] if f.get('featureId') in iso.get('sequenceIds', [])]
        edit_records = []
        for f in sorted(edits, key=lambda x: x['location']['start']['value'], reverse=True):
            start, end = f['location']['start']['value'], f['location']['end']['value']
            alternative = f.get('alternativeSequence', {})
            original = alternative.get('originalSequence')
            if original and seq[start - 1:end] != original:
                raise ValueError('Official edit does not match parent sequence: ' + acc)
            replacement = alternative.get('alternativeSequences', [''])[0]
            seq = seq[:start - 1] + replacement + seq[end:]
            edit_records.append({'feature_id': f['featureId'], 'canonical_start': start,
                                 'canonical_end': end, 'replacement': replacement})
        fasta = ''.join(line.strip() for line in path.read_text().splitlines() if not line.startswith('>'))
        if seq != fasta:
            raise ValueError('Official alternative sequence differs from FASTA: ' + acc)
        seqs[acc] = seq
        variants.append({'accession': acc, 'official_isoform': iso, 'edits': edit_records,
                         'length': len(seq), 'sequence_sha256': hashlib.sha256(seq.encode()).hexdigest(),
                         'matches_official_fasta': True})
    return records, seqs, variants


def interaction_records(paths, selection, official, sequences):
    with paths['intact'].open() as stream:
        lines = list(csv.reader(stream, delimiter='\t'))

    def has_id(fields, accession, side):
        identifiers = (fields[side] + '|' + fields[side + 2]).lower().split('|')
        return 'uniprotkb:' + accession.lower() in identifiers

    joins = []
    for item in selection['bindings']:
        partners = []
        for partner in item['partners']:
            accession = partner.split(':', 1)[1]
            parent = official[accession.split('-')[0]]
            taxon = str(parent['organism']['taxonId'])
            rogid = base64.b64encode(hashlib.sha1(sequences[accession].encode()).digest()).decode().rstrip('=') + taxon
            hits = []
            for lineno, fields in enumerate(lines, 1):
                if 'pubmed:' + item['reference'].split(':')[-1] not in fields[8].split('|'):
                    continue
                sides = [s for s in (0, 1) if has_id(fields, selection['target'], s) and has_id(fields, accession, 1 - s)]
                for side in sides:
                    other = 1 - side
                    hits.append({'line_1based': lineno, 'target_side': side, 'raw_fields': fields,
                                 'partner_reference_checksum': fields[32 + other], 'expected_ROGID': rogid,
                                 'reference_sequence_match': 'rogid:' + rogid in fields[32 + other].split('|'),
                                 'target_features': fields[36 + side], 'partner_features': fields[36 + other]})
            partners.append({'partner': partner, 'records': hits})
        joins.append({'source_row': item['source_row'], 'source_reference': item['reference'], 'partners': partners})
    return {'total_records': len(lines), 'joins': joins}


def phylogeny(paths, selection, official):
    tree = json.loads(paths['tree'].read_text())['search']['tree_topology']['annotation_node']
    nodes, lineage = {}, {}

    def visit(node, ancestors):
        key = node['persistent_id']
        nodes[key] = {k: v for k, v in node.items() if k != 'children'}
        lineage[key] = ancestors + [key]
        children = node.get('children', {}).get('annotation_node', [])
        if isinstance(children, dict):
            children = [children]
        for child in children:
            visit(child, lineage[key])

    visit(tree, [])
    target = [k for k, v in nodes.items() if 'UniProtKB=' + selection['target'] in v.get('node_name', '').split('|')]
    if len(target) != 1:
        raise ValueError('Target absent or nonunique in tree')
    target = target[0]
    msa = json.loads(paths['msa'].read_text())['search']['MSA_list']['sequence_info']
    selected = [v for v in msa if v.get('persistent_id') == target]
    if len(selected) != 1:
        raise ValueError('Target absent or nonunique in MSA')
    sequence = selected[0]['sequence'].replace('.', '').replace('-', '').upper()
    if sequence != official[selection['target']]['sequence']['value']:
        raise ValueError('MSA target does not match official sequence')
    wanted = {(x['node'], x['go']) for x in selection['paint_assertions']}
    assertions = []
    with paths['ibd'].open() as stream:
        for lineno, line in enumerate(stream, 1):
            if line.startswith('!'):
                continue
            fields = line.rstrip('\n').split('\t')
            if len(fields) > 7 and (fields[1], fields[4]) in wanted and fields[6] == 'IBD':
                assertions.append({'line_1based': lineno, 'fields': fields,
                                   'on_target_path': fields[1] in lineage[target]})
    donors = sorted({v for row in assertions for v in row['fields'][7].split('|')})
    donor_paths = []
    for donor in donors:
        token = donor.replace(':', '=')
        if donor.startswith('WB:'):
            token = 'WormBase=' + donor.split(':', 1)[1]
        matches = [k for k, v in nodes.items() if token in v.get('node_name', '').split('|') or v.get('gene_id') == donor]
        donor_paths.append({'donor': donor, 'matches': [{'node': nodes[k], 'path': lineage[k]} for k in matches]})
    return {'nodes': len(nodes), 'msa_records': len(msa), 'target_node': nodes[target],
            'target_path': [nodes[k] for k in lineage[target]], 'target_sequence_length': len(sequence),
            'sequence_exact': True, 'selected_IBDs': assertions, 'donor_paths': donor_paths}


def reactome_records(paths):
    graph = {k[9:]: json.loads(p.read_text()) for k, p in paths.items() if k.startswith('reactome:')}
    allowed = ['dbId', 'stId', 'displayName', 'schemaClass', 'speciesName', 'input', 'output',
               'hasMember', 'hasComponent', 'referenceEntity', 'compartment', 'startCoordinate', 'endCoordinate']
    return {k: {field: value for field, value in v.items() if field in allowed} for k, v in graph.items()}


def workbook_records(paths, selection):
    import openpyxl
    import xlrd
    needles = {selection['target'].upper(), selection['symbol'].upper()}

    def matches(cells):
        return any(needle in str(cell).upper() for cell in cells for needle in needles)

    parotid = []
    book = xlrd.open_workbook(paths['parotid_table'])
    for sheet in book.sheets():
        hits = [{'row_1based': row + 1, 'cells': sheet.row_values(row)[:10]}
                for row in range(sheet.nrows) if matches(sheet.row_values(row)[:10])]
        parotid.append({'sheet': sheet.name, 'rows': sheet.nrows, 'columns': sheet.ncols,
                        'selected_columns_1based': list(range(1, 11)),
                        'headers': sheet.row_values(0)[:10], 'hits': hits})
    matrix = []
    book = openpyxl.load_workbook(paths['matrix_table'], read_only=True, data_only=True)
    for sheet in book:
        rows = list(sheet.iter_rows(values_only=True))
        hits = [{'row_1based': index + 1, 'cells': list(row)} for index, row in enumerate(rows) if matches(row)]
        # Keep short column headings, not the article's caption prose.
        headers = [{'row_1based': index + 1, 'cells': list(row)} for index, row in enumerate(rows[:6])
                   if any(cell is not None for cell in row) and not any(isinstance(cell, str) and len(cell) > 100 for cell in row)]
        matrix.append({'sheet': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                       'headers': headers, 'hits': hits})
    book.close()
    return {'PMID:19199708': parotid, 'PMID:28327460': matrix}


def ontology_records(paths, selection):
    terms = {}
    for key, path in paths.items():
        if not key.startswith('ontology:'):
            continue
        response = json.loads(path.read_text())
        for term in response.get('results', []):
            if term.get('id') in selection['ontology_terms']:
                terms.setdefault(term['id'], []).append({'input_id': key, 'term': term})
    missing = set(selection['ontology_terms']) - set(terms)
    if missing:
        raise ValueError('Missing selected term definitions: ' + str(sorted(missing)))
    return terms


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inputs', type=Path, help='JSON mapping input IDs to saved raw paths')
    parser.add_argument('--download-cache', type=Path, help='Download missing raw inputs from recorded public URLs')
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--selection', type=Path, default=Path(__file__).with_name('selection.json'))
    args = parser.parse_args()
    manifest = json.loads(Path(__file__).with_name('source-manifest.json').read_text())
    mapping = json.loads(args.inputs.read_text()) if args.inputs else {}
    paths = load_inputs(manifest, mapping, args.download_cache)
    selection = json.loads(args.selection.read_text())
    official, sequences, products = official_sequences(paths)
    interactions = interaction_records(paths, selection, official, sequences)
    partner_data = {acc: {'names': data.get('proteinDescription'), 'genes': data.get('genes'),
                         'organism': data.get('organism'), 'sequence_length': data['sequence']['length'],
                         'features': [f for f in data.get('features', []) if f['type'] in
                                      ['Domain', 'Active site', 'Binding site', 'Motif', 'Coiled coil', 'Signal',
                                       'Propeptide', 'Chain', 'Peptide', 'Alternative sequence', 'Region', 'Transmembrane']],
                         'comments': [c for c in data.get('comments', []) if c['commentType'] in ['FUNCTION', 'CATALYTIC ACTIVITY', 'CAUTION', 'SIMILARITY', 'ALTERNATIVE PRODUCTS']],
                         'GO': [r for r in data.get('uniProtKBCrossReferences', []) if r['database'] == 'GO']}
                    for acc, data in official.items()}
    outputs = {'interactions.json': interactions, 'official-partners.json': partner_data,
               'alternative-sequences.json': products, 'phylogeny.json': phylogeny(paths, selection, official),
               'reactome.json': reactome_records(paths), 'workbooks.json': workbook_records(paths, selection),
               'ontology.json': ontology_records(paths, selection)}
    args.output.mkdir(parents=True, exist_ok=False)
    for filename, value in outputs.items():
        (args.output / filename).write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps({'inputs': len(paths), 'binding_sources': len(interactions['joins']),
                      'interaction_records': interactions['total_records'], 'official_products_reconstructed': len(products)}))


if __name__ == '__main__':
    main()
