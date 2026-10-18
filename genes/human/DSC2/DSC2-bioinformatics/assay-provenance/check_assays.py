"""Recompute reference identities and published-workbook provenance from real inputs."""
from pathlib import Path
import argparse,base64,hashlib,json
import openpyxl

def analyze(d):
    records={a:json.loads((d/f'UniProt-{a}.json').read_text()) for a in ['Q02487','Q14126']}
    sequences={a:r['sequence']['value'] for a,r in records.items()}
    for a,r in records.items():
        if r['primaryAccession']!=a or r['organism']['taxonId']!=9606:
            raise ValueError(f'Unexpected reference identity: {a}')
    rogids={a:base64.b64encode(hashlib.sha1(s.encode()).digest()).decode().rstrip('=')+'9606' for a,s in sequences.items()}
    pairs=[]
    for n,line in enumerate((d/'IntAct-Q02487.mitab').read_text().splitlines()):
        fields=line.split('\t')
        if len(fields)>38 and 'pubmed:21220045' in fields[8] and set(fields[:2])=={'uniprotkb:Q02487','uniprotkb:Q14126'}:
            verified=all(fields[32+i]=='rogid:'+rogids[fields[i].split(':')[1]] for i in [0,1])
            pairs.append({'zero_based_record':n,'participants':fields[:2],'detection_method':fields[6],'source':fields[8],'interaction_type':fields[11],'stable_ids':fields[13],'experimental_host':fields[28],'features':fields[36:38],'reference_rogids_match':verified})
    w=openpyxl.load_workbook(d/'41586_2022_5028_MOESM4_ESM.xlsx',read_only=True,data_only=True)
    rows=list(w['Table 1'].iter_rows(values_only=True));header=rows[5]
    targets=[{'excel_row':i+1,'record':dict(zip(header,row))}for i,row in enumerate(rows)if row[0]=='Q02487']
    for target in targets:
        seq=target['record']['AA_seq'];full=sequences['Q02487']
        target['exact_reference_start_positions']=[i+1 for i in range(len(full))if full.startswith(seq,i)]
        target['length']=len(seq)
    rows=list(w['Table 3'].iter_rows(values_only=True));header=rows[6]
    affinities=[{'excel_row':i+1,'record':{k:v for k,v in zip(header,row)if k is not None}}for i,row in enumerate(rows)if len(row)>2 and set(row[1:3])=={'Q02487','Q14126'}]
    w=openpyxl.load_workbook(d/'41586_2022_5028_MOESM6_ESM.xlsx',read_only=True,data_only=True)
    networks={}
    for sheet in ['Figure1_network','Figure2_network']:
        rows=list(w[sheet].iter_rows(values_only=True));header=rows[0]
        networks[sheet]=[{'excel_row':i+1,'record':dict(zip(header,row))}for i,row in enumerate(rows[1:],1)if 'Q02487'in row and 'Q14126'in row]
    return {'inputs':{p.name:hashlib.sha256(p.read_bytes()).hexdigest()for p in sorted(d.iterdir())if p.is_file()},'reference_lengths':{a:len(s)for a,s in sequences.items()},'reference_rogids':rogids,'intact_pairs':pairs,'library_targets':targets,'curated_affinities':affinities,'network_records':networks,'limits':['Reference ROGIDs do not authenticate every mutant experimental insert.','A shared ectodomain cannot distinguish natural cytoplasmic tails.','The workbook Source and References fields determine whether affinity was measured here or imported from prior literature.','A low screen score is not a measurement of zero binding affinity.']}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--inputs',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args()
    result=analyze(a.inputs)
    if a.output.exists():raise FileExistsError(a.output)
    a.output.write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps({'target_records':len(result['library_targets']),'intact_pairs':len(result['intact_pairs']),'reference_lengths':result['reference_lengths']}))
